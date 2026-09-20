"""文档解析:把 pdf/docx/doc/xlsx/txt 统一解析成纯文本。

用法:
    text = parse_document(Path("xxx.pdf"))
"""
from __future__ import annotations

import subprocess
import sys
import tempfile
from pathlib import Path

import openpyxl
from src.utils.stderr_log import warn as _warn


def parse_document(path: Path) -> str:
    """根据扩展名分派。失败返回空字符串,不抛异常。"""
    ext = path.suffix.lower()
    try:
        if ext == ".pdf":
            return _parse_pdf(path)
        if ext == ".docx":
            return _parse_docx(path)
        if ext == ".doc":
            return _parse_doc(path)
        if ext in (".xlsx", ".xlsm"):
            return _parse_xlsx(path)
        if ext == ".pptx":
            return _parse_pptx(path)
        if ext in (".txt", ".md"):
            return path.read_text(encoding="utf-8", errors="ignore")
    except Exception as e:  # noqa: BLE001
        _warn(f"[doc_parser] 解析失败 {path.name}: {e}")
    return ""


def _parse_pdf(path: Path) -> str:
    import pdfplumber
    parts: list[str] = []
    with pdfplumber.open(path) as pdf:
        for i, page in enumerate(pdf.pages, 1):
            text = page.extract_text() or ""
            if text.strip():
                parts.append(f"[第{i}页]\n{text}")
    return "\n\n".join(parts)


def _parse_docx(path: Path) -> str:
    import docx  # python-docx
    d = docx.Document(str(path))
    paragraphs = [p.text for p in d.paragraphs if p.text.strip()]
    # 表格内容也拼进来
    for tbl in d.tables:
        for row in tbl.rows:
            paragraphs.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(paragraphs)


def _parse_doc(path: Path) -> str:
    """老式 .doc:先调 libreoffice 转成 .docx,再用 python-docx 读。"""
    with tempfile.TemporaryDirectory() as tmpdir:
        cmd = [
            "libreoffice", "--headless",
            "--convert-to", "docx",
            "--outdir", tmpdir,
            str(path),
        ]
        try:
            subprocess.run(cmd, check=True, capture_output=True, timeout=60)
        except (FileNotFoundError, subprocess.CalledProcessError) as e:
            _warn(f"[doc_parser] libreoffice 转换失败 {path.name}: {e}. "
                  f"请确认已安装: sudo apt install libreoffice-core")
            return ""
        converted = Path(tmpdir) / (path.stem + ".docx")
        if not converted.exists():
            return ""
        return _parse_docx(converted)


def _parse_xlsx(path: Path) -> str:
    # 【2026-09-20 修】这里原来是 `read_only=True`。openpyxl 在只读模式下**只信
    # sheet XML 里的 `<dimension>` 标签**,而且**不校验它对不对**:成绩记分册里
    # 写的是 `<dimension ref="A1"/>`(错的范围声明,它其实有 146 行),于是这份
    # 140 个学生的工作簿被读成 1 行 1 列,**不报错** —— 进索引时就只贡献了标题那一格。
    # 语料静默缩水的形状和检索退化一模一样,查起来极难 —— 所以宁可用普通模式,
    # 多花一点内存换"解析结果可信"。见 knowledge_base.md §8 第 12 条。
    wb = openpyxl.load_workbook(str(path), data_only=True, read_only=False)
    parts: list[str] = []
    for sheet in wb.sheetnames:
        ws = wb[sheet]
        lines: list[str] = [f"[Sheet: {sheet}]"]
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                lines.append(" | ".join(cells))
        parts.append("\n".join(lines))
    return "\n\n".join(parts)


def _parse_pptx(path: Path) -> str:
    """PPTX 按页取文字。

    为什么**不用** python-pptx 的 `shape.text_frame.text` 直接遍历顶层形状:
    课件里的要点常常被作者分到"组合形状"(group)里,组合内的文字在顶层遍历时
    **一个都取不到**(实测张老师的错误集锦 PPT 有大量组合,漏掉就只剩页标题)。
    所以这里递归展开组合,并且把表格按行拼成文本 —— 错误集锦里的"错误/正确"
    对照就是用表格排的。
    """
    from pptx import Presentation
    parts: list[str] = []
    prs = Presentation(str(path))
    for i, slide in enumerate(prs.slides, 1):
        lines = [ln for ln in _pptx_shape_text(slide.shapes) if ln]
        # 备注页里也常有讲解文字,一并带上
        try:
            if slide.has_notes_slide:
                note = (slide.notes_slide.notes_text_frame.text or "").strip()
                if note:
                    lines.append(f"(备注){note}")
        except Exception:  # noqa: BLE001
            pass
        if lines:
            parts.append(f"[第{i}页]\n" + "\n".join(lines))
    return "\n\n".join(parts)


def _pptx_shape_text(shapes) -> list[str]:
    """递归收集形状里的文字(含组合形状与表格)。"""
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    out: list[str] = []
    for sh in shapes:
        try:
            if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
                out.extend(_pptx_shape_text(sh.shapes))
                continue
            if getattr(sh, "has_text_frame", False):
                t = (sh.text_frame.text or "").strip()
                if t:
                    out.append(t)
            if getattr(sh, "has_table", False):
                for row in sh.table.rows:
                    cells = [c.text.strip() for c in row.cells]
                    if any(cells):
                        out.append(" | ".join(cells))
        except Exception:  # noqa: BLE001
            continue
    return out
