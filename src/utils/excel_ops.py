"""Excel 操作:花名册读取、补交表读写。"""
from __future__ import annotations

import re
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl

from src import config
from src.utils.stderr_log import warn as _warn

# 学号:6 位以上纯数字。**不写死 8 位** —— 写死了的话,哪天教务导出 10 位学号,
# 这份名单会被安静地读成一个空字典,和这个模块原来那个 bug 一模一样的形状。
_ID_RE = re.compile(r"^\d{6,}$")


def _cell_text(v: object) -> str:
    """把单元格值规范化成字符串。

    学号在不同导出里可能是文本、也可能是数字 —— openpyxl 把后者读成 int,
    位数一多就读成 float(2530123456.0),`str()` 出来带个 ".0",学号就永远对不上。
    这是"看着像读到了、其实一个都对不上"的那类坑,所以在入口处一次性抹平。
    """
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def _load(path: Path):
    """打开工作簿。

    【2026-09-20 修】原来是 `read_only=True`,而 openpyxl 在只读模式下**只信
    sheet XML 里的 `<dimension>` 标签**,且**不校验它对不对**。成绩记分册里写的是
    `<dimension ref="A1"/>` —— 一个错的范围声明(它明明有 146 行 17 列),
    于是被读成 **1 行 1 列**:不抛异常、不告警,只是安静地什么都读不到。
    它和"路径配错"叠加在一起,让"学号/姓名校验"从未生效过
    (见 knowledge_base.md §8 第 12 条,`test_roster.py` 里钉了它的复现)。

    花名册只有一两百行,普通模式的内存代价可以忽略,换来的是"读到的就是真数据"。
    """
    return openpyxl.load_workbook(str(path), data_only=True, read_only=False)


@dataclass
class RosterScan:
    """一次花名册读取的**全过程**,包括失败原因。

    为什么要返回结构体而不是直接返回 dict:这个模块出过的两个 bug(路径配错、
    read_only 读成 1×1)**都是静默的** —— 上层只看到"空名单",看不到为什么空。
    `scripts/check_roster.py` 拿它打印诊断,把"静默降级"变成"跑一次就知道"。
    """
    path: Path
    ok: bool = False
    reason: str = ""
    header_row: int = 0
    id_col: int = 0
    name_col: int = 0
    skipped_rows: int = 0          # 有内容但学号/姓名不合规的行
    duplicates: int = 0            # 同学号不同姓名,保留先出现的
    roster: dict[str, str] = field(default_factory=dict)

    @property
    def count(self) -> int:
        return len(self.roster)

    @property
    def id_lengths(self) -> dict[int, int]:
        return dict(Counter(len(k) for k in self.roster))


def scan_roster(path: Path | None = None) -> RosterScan:
    """读花名册,连"为什么没读出来"一起返回。**不抛异常。**"""
    p = path or config.GRADEBOOK_PATH
    if not p.exists():
        return RosterScan(p, reason=f"文件不存在: {p}")

    try:
        wb = _load(p)
        ws = wb.active
    except Exception as e:  # noqa: BLE001 —— 坏文件不该让补交功能整个挂掉
        return RosterScan(p, reason=f"打不开: {type(e).__name__}: {e}")

    # 找表头:哪一行同时出现了学号类表头和姓名类表头。
    # 不按固定行号找,是因为花名册每学期都换(雨课堂导出、教务导出、记分册…),
    # 版式一变,写死的行号就会安静地读成空名单 —— 那正是这个 bug 藏了几个月的原因。
    hit: tuple[int, int, int] | None = None
    for r in range(1, min(ws.max_row, config.ROSTER_HEADER_SCAN_ROWS) + 1):
        vals = [_cell_text(ws.cell(r, c).value) for c in range(1, ws.max_column + 1)]
        id_col = next((i + 1 for i, v in enumerate(vals)
                       if v in config.ROSTER_ID_HEADERS), 0)
        name_col = next((i + 1 for i, v in enumerate(vals)
                         if v in config.ROSTER_NAME_HEADERS), 0)
        if id_col and name_col:
            hit = (r, id_col, name_col)
            break

    if hit is None:
        return RosterScan(
            p,
            reason=(f"前 {config.ROSTER_HEADER_SCAN_ROWS} 行里没有任何一行同时含 "
                    f"{'/'.join(config.ROSTER_ID_HEADERS)} 和 "
                    f"{'/'.join(config.ROSTER_NAME_HEADERS)} 表头"),
        )

    header_row, id_col, name_col = hit
    scan = RosterScan(p, header_row=header_row, id_col=id_col, name_col=name_col)
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        sid = _cell_text(row[id_col - 1]) if len(row) >= id_col else ""
        name = _cell_text(row[name_col - 1]) if len(row) >= name_col else ""
        if not sid and not name:
            continue                      # 空行,不算异常
        if not _ID_RE.match(sid) or not name:
            scan.skipped_rows += 1
            continue
        if sid in scan.roster:
            if scan.roster[sid] != name:
                scan.duplicates += 1
            continue
        scan.roster[sid] = name

    scan.ok = True
    return scan


def load_roster(path: Path | None = None) -> dict[str, str]:
    """读花名册 → {学号: 姓名}。定位方式见 scan_roster()。

    **找不到/读不出时返回空字典,不抛异常** —— 补交登记不该因为名单读不出来就整个失败。
    但失败**一定发出告警**:这个模块原先的两个 bug 都是"安静地返回空字典",
    于是"学号校验"停摆了好几个月没人发现。成功时不打日志(微信链路上每次补交都会调它)。
    """
    scan = scan_roster(path)
    if not scan.ok:
        _warn(f"[excel_ops] 花名册读不出来 —— {scan.reason}。"
              f"补交校验会退化成「名单缺失」。")
        return {}
    if not scan.roster:
        _warn(f"[excel_ops] 花名册打开了、表头也找到了(第 {scan.header_row} 行),"
              f"但**一个学生都没解析出来**(跳过 {scan.skipped_rows} 行)。"
              f"检查学号列是不是真的数字:{scan.path}")
        return {}
    return scan.roster


def ensure_submission_table(path: Path) -> None:
    """补交表不存在则新建一张带表头的空表。

    **自动新建是对的**:换学期就是换一门课,新学期第一笔补交理应把表建出来
    (路径按课推导,见 config.submission_table_for)。

    **但父目录不存在时不许建。** 补交表的位置是"花名册旁边",父目录(花名册/)
    不在就说明路径本身是错的 —— 这时候 mkdir + 建空表,等于在一个助教根本不会
    去看的地方悄悄记下补交,而机器人照样回「✅ 已登记」。这类"静默分裂"在 FAQ 那边
    已经定过规矩(见 src/handlers/record.py:路径可疑时宁可不写)。所以这里直接抛错,
    由调用方翻成一句诚实的回执。
    """
    if path.exists():
        return
    if not path.parent.exists():
        raise FileNotFoundError(
            f"补交表所在目录不存在,不敢新建:{path.parent}"
            f"(补交表应当与花名册同目录,检查 config.GRADEBOOK_PATH)"
        )
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "补交记录"
    ws.append(config.SUBMISSION_HEADERS)
    # 简单格式化:表头加粗
    from openpyxl.styles import Font
    bold = Font(bold=True)
    for cell in ws[1]:
        cell.font = bold
    # 列宽
    widths = {"A": 18, "B": 12, "C": 10, "D": 10, "E": 60, "F": 18}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    wb.save(str(path))
    _warn(f"[excel_ops] 新建补交表: {path}")


def append_submission_row(path: Path, row: dict) -> None:
    """追加一行补交记录。读-改-存模式。"""
    ensure_submission_table(path)
    wb = openpyxl.load_workbook(str(path))
    ws = wb.active
    ws.append([row.get(h, "") for h in config.SUBMISSION_HEADERS])
    wb.save(str(path))


def load_submission_history(path: Path | None = None) -> list[dict]:
    """读出补交表全部历史记录(用于重复检测、统计)。"""
    p = path or config.SUBMISSION_TABLE_PATH
    if not p.exists():
        return []
    # 和花名册同一个理由不用 read_only:见 _load() 的注释
    wb = _load(p)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    headers = list(rows[0])
    return [dict(zip(headers, r)) for r in rows[1:] if any(r)]
